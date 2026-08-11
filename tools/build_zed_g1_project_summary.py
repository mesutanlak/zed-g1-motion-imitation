from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "ZED2i_G1_EDU_Proje_Ozeti.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEXT = "1F2933"
MUTED = "5B6573"
LIGHT = "F2F4F7"
PALE_BLUE = "EAF2F8"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_paragraph_spacing(p, before=0, after=6, line=1.10):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_run(p, text, bold=False, color=TEXT, size=9.7, italic=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def add_body(doc, text, after=3.4):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after, line=1.02)
    add_run(p, text)
    return p


def add_bullet(doc, lead, text, after=2.7):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, after=after, line=1.02)
    add_run(p, lead, bold=True, color=DARK_BLUE)
    add_run(p, text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    return p


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(paragraph, "ZED 2i × G1 EDU  |  ", color=MUTED, size=8.5)
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    run.font.name = "Calibri"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_status_box(doc):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.45)
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_BLUE)
    set_cell_margins(cell, top=105, bottom=105, start=150, end=150)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=0, line=1.03)
    add_run(p, "MEVCUT DURUM  ", bold=True, color=DARK_BLUE, size=9.5)
    add_run(
        p,
        "Tek ZED 2i’den canlı BODY_38 alımı, çevrimdışı kalite analizi, GMR tabanlı "
        "23-DOF referans üretimi ve MuJoCo/Isaac Lab simülasyon zinciri çalışır durumdadır. "
        "Fiziksel G1’e motor komutu veren DDS çıkışı bilinçli olarak kapalıdır.",
        size=9.5,
    )


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.58)
sec.bottom_margin = Inches(0.58)
sec.left_margin = Inches(1.0)
sec.right_margin = Inches(1.0)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9.7)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(3.4)
normal.paragraph_format.line_spacing = 1.02

for name, size, color, before, after in (
    ("Heading 1", 13.5, BLUE, 7, 3),
    ("Heading 2", 11.8, BLUE, 6, 3),
    ("Heading 3", 10.8, DARK_BLUE, 5, 2),
):
    st = styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

list_style = styles["List Bullet"]
list_style.font.name = "Calibri"
list_style.font.size = Pt(9.5)
list_style.paragraph_format.left_indent = Inches(0.5)
list_style.paragraph_format.first_line_indent = Inches(-0.25)

header = sec.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_paragraph_spacing(hp, after=2, line=1.0)
add_run(hp, "TEKNİK PROJE ÖZETİ  •  29 TEMMUZ 2026", bold=True, color=MUTED, size=8.5)
p_pr = hp._p.get_or_add_pPr()
p_bdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "9")
bottom.set(qn("w:space"), "4")
bottom.set(qn("w:color"), BLUE)
p_bdr.append(bottom)
p_pr.append(p_bdr)

add_page_number(sec.footer.paragraphs[0])

title = doc.add_paragraph()
set_paragraph_spacing(title, before=2, after=1, line=1.0)
add_run(title, "ZED 2i ile G1 EDU İnsan Hareketi Taklit Sistemi", bold=True, color=DARK_BLUE, size=17.5)
subtitle = doc.add_paragraph()
set_paragraph_spacing(subtitle, after=8, line=1.0)
add_run(
    subtitle,
    "Veri toplama, kalite analizi, hareket retargeting’i, fizik tabanlı simülasyon ve güvenli sim2real yol haritası",
    color=MUTED,
    size=9.8,
    italic=True,
)

add_status_box(doc)

add_heading(doc, "1. Amaç ve sistemin genel yapısı", 1)
add_body(
    doc,
    "Projenin amacı, ZED 2i stereo derinlik kamerasının gördüğü bir insanın yerinde yaptığı "
    "gövde, kol ve bacak hareketlerini Unitree G1 EDU 23-DOF humanoid robota aktarabilen; "
    "önce simülasyonda doğrulanan, daha sonra güvenlik kapılarıyla fiziksel robota taşınan bir "
    "taklit sistemi geliştirmektir. ZED’in iki lensi ayrı iki kamera gibi bağımsız iskelet üretmez; "
    "stereo çift, tek bir ZED cihazı içinde derinlik ve 3B konum üretir. Uçtan uca akış "
    "“ZED BODY_38 → kalite kapısı → insan/robot koordinat dönüşümü → GMR ters kinematik → "
    "G1 23-DOF hedefleri → MuJoCo/Isaac Lab” biçimindedir.",
)

add_heading(doc, "2. Kullanılan metodoloji ve proje kodları", 1)
add_bullet(
    doc,
    "Algılama ve kayıt: ",
    "zed_g1_skeleton.py, Stereolabs ZED SDK 5.4 BODY_38 modelinden 38 adet 3B keypoint, "
    "nokta güvenleri, pelvis/root pozu, yerel quaternionlar, 2B pikseller, kamera iç/dış "
    "parametreleri, kişinin 3B uzaklığı ve ZED 2i IMU verisini alır. Kayıtlar kare-zamanlı JSONL; "
    "stereo görüntü ve yeniden işleme kaynağı ise isteğe bağlı SVO2 olarak saklanır.",
)
add_bullet(
    doc,
    "Kalite analizi: ",
    "analyze_body38_recording.py ve analyze_recording_collection.py; etkin FPS, kare düşmesi, "
    "izleme sürekliliği, gövde/kol/bacak/ayak görünürlüğü, kemik uzunluğu tutarlılığı, quaternion "
    "normu, keypoint hızı ve 2–4 m çalışma mesafesini ölçerek kayıtları “whole_body_candidate”, "
    "“upper_body_only” veya “partial_or_retake” olarak sınıflandırır.",
)
add_bullet(
    doc,
    "Retargeting: ",
    "body38_to_gmr.py ZED iskeletini pelvis merkezli, zemin hizalı ve kısa kayıplarda bellekle "
    "tamamlanan GMR insan verisine çevirir. YanjieZe/GMR, resmi G1 kinematiği üzerinde IK ile "
    "29-DOF çözüm üretir; g1_dof_projection.py fiziksel 23-DOF sürümde bulunmayan altı ekseni "
    "isimle çıkarır. gmr_live_bridge.py 7 Hz süzme, eklem hız sınırı ve 0,35 s eski-veri "
    "watchdog’u sonrası 23 hedefi UDP 15051’e yayınlar.",
)
add_bullet(
    doc,
    "Simülasyon: ",
    "g1_zed_live_mimic.py resmi Unitree MuJoCo modelinde yerçekimi, temas, ayak-zemin kontrolü "
    "ve düşüş resetiyle hızlı sim2sim sınaması sağlar. isaac_g1_23dof_live.py ise ayrı Windows "
    "Isaac Sim 5.0 / Isaac Lab v2.2 ortamında resmi URDF ve aktüatör tork-hız sınırlarını kullanır. "
    "RViz, JointState/TF ve daha önce kurulan LiDAR–derinlik sensörü hattının görsel tanısı için korunmuştur.",
)

add_heading(doc, "3. Toplanan veriler ve ölçülen çıktılar", 1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Inches(1.55), Inches(2.15), Inches(2.75)]
for idx, width in enumerate(widths):
    table.columns[idx].width = width
hdr = table.rows[0].cells
for idx, text_value in enumerate(("Çıktı", "Ölçüm", "Yorum")):
    set_cell_shading(hdr[idx], BLUE)
    set_cell_margins(hdr[idx])
    hdr[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = hdr[idx].paragraphs[0]
    set_paragraph_spacing(p, after=0, line=1.0)
    add_run(p, text_value, bold=True, color=WHITE, size=9.3)
set_repeat_table_header(table.rows[0])

rows = [
    (
        "Ana başlangıç kaydı\n134332",
        "5.268 kare; 176,3 s; 29,86 FPS; %96,95 body confidence",
        "%90,83 tüm-vücut çekirdeği; veri kümesi/retargeting için en güçlü başlangıç.",
    ),
    (
        "Yakın dönem kaydı\n105156",
        "573 kare; 28,60 FPS; %81,68 üst gövde; %74,35 tüm vücut",
        "Kollar izlenebilir; tek ön kamerada ayak desteği ve örtüşen uzuvlar zayıf.",
    ),
    (
        "GMR + Isaac Lab",
        "23 eklem paketi; testte 12 geçerli hedef; GPU PhysX 100 adım",
        "Canlı zincir doğrulandı; denge politikası ve sürücü uyumlu GUI halen geliştirme konusu.",
    ),
]
for row_index, row_values in enumerate(rows):
    cells = table.add_row().cells
    for idx, value in enumerate(row_values):
        set_cell_shading(cells[idx], LIGHT if row_index % 2 == 0 else WHITE)
        set_cell_margins(cells[idx])
        p = cells[idx].paragraphs[0]
        set_paragraph_spacing(p, after=0, line=1.0)
        add_run(p, value, size=8.5, bold=(idx == 0), color=DARK_BLUE if idx == 0 else TEXT)

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

add_heading(doc, "4. Bulgular, sınırlar ve güvenlik yaklaşımı", 1)
add_body(
    doc,
    "Tek ZED 2i, omuz–dirsek–bilek ve temel kalça–diz–ayak bileği hedeflerini üretmek için "
    "yeterli bir başlangıçtır; ancak önden görünüşte bacağın diğer bacağı örtmesi, ayakların küçük "
    "görünmesi ve ellerin gövde önünde kalması doğruluğu düşürür. BODY_38, Dex3 üç parmaklı elin "
    "motorlarını benzersiz çözmeye yetecek falanks ayrıntısı sağlamaz; SVO2 görüntülerinden ayrı "
    "21-landmark el takibi ve dex-retargeting gerekir. Görünür iskelet doğruluğu da motion-capture "
    "ground truth değildir; kalibre edilmiş referans hareketlerle hata ölçümü yapılmalıdır.",
)
add_bullet(
    doc,
    "Simülasyon güvenliği: ",
    "Varsayılan upper_body modunda bacaklar nominal duruşta kalır; whole_body ancak ayak teması "
    "ve denge politikası tamamlandığında açılır. Eklem limitleri, hız kırpma, zaman aşımı, nötr "
    "poza yumuşak dönüş, pelvis yüksekliği/dikliği denetimi ve fizik reseti uygulanır.",
)
add_bullet(
    doc,
    "Fiziksel robot güvenliği: ",
    "Mevcut paketlerde physical_robot_output=false olup Unitree DDS motor çıkışı yoktur. Gerçek "
    "G1 deneyi; askı, düşük hız/tork, mekanik çalışma alanı ve self-collision sınırları, destek "
    "poligonu/denge denetimi, dead-man, watchdog ve operatör acil durdurması olmadan başlatılmamalıdır.",
)
add_bullet(
    doc,
    "Teknik kısıt: ",
    "RTX 5090 ile Isaac Lab başsız GPU PhysX testi çalışmıştır; mevcut NVIDIA 610.62 sürücüsünde "
    "Isaac Sim GUI device-lost sorunu görülmüştür. Bu, veri hattından bağımsız bir grafik sürücüsü "
    "uyumluluk problemidir ve doğrulanmış sürücüyle yeniden sınanmalıdır.",
)

add_heading(doc, "5. Gelecekte iki ZED 2i ile geliştirme", 1)
add_body(
    doc,
    "İkinci ZED 2i, tek cihazın stereo çiftini çoğaltmak için değil, farklı bakış açısından "
    "örtüşmeleri azaltmak için eklenmelidir. Kameralar yaklaşık 60–90° ayrık açıyla tüm vücudu "
    "ortak görecek biçimde yerleştirilir; seri numarasıyla sabitlenir, ZED360 ile extrinsic "
    "kalibre edilir ve Stereolabs Fusion API üzerinden aynı WORLD koordinatında birleştirilir. "
    "Kareler SDK zaman damgasıyla eşleştirilir; aynı kişi gövde geometrisi/ID ile ilişkilendirilir; "
    "her keypoint için güven-ağırlıklı seçim veya füzyon yapılarak tek bir BODY_38 akışı üretilir. "
    "Böylece yandan görülen diz, ayak ve gövde önündeki el diğer kamera tarafından tamamlanabilir.",
)
add_body(
    doc,
    "İki kameranın aynı USB denetleyicisine yük bindirmemesi için doğrudan USB 3.x portları veya "
    "bağımsız PCIe USB denetleyicileri kullanılmalı; powered hub’ın yalnızca güç sağladığı, bant "
    "genişliğini artırmadığı dikkate alınmalıdır. Önce eşzamanlı JSONL+SVO2 kaydı alınmalı, tek ve "
    "çift kamera aynı hareketlerde görünürlük, süreklilik ve retargeting hatasıyla A/B karşılaştırılmalıdır. "
    "Fusion çıktısı mevcut GMR/23-DOF sözleşmesini koruduğundan MuJoCo ve Isaac Lab tarafı değişmeden kullanılabilir.",
)

add_heading(doc, "6. Önerilen geliştirme yol haritası", 1)
add_bullet(
    doc,
    "1 — Veri standardizasyonu: ",
    "A/T nötr pozu, yavaş tek-eklem hareketleri, çömelme/tek ayak ve örtüşme senaryolarını etiketli "
    "kliplere ayır; düşük kaliteli segmentleri eğitimden çıkar.",
)
add_bullet(
    doc,
    "2 — Öğrenilen takip politikası: ",
    "GMR q_ref[23], hızlar ve ayak temaslarını Isaac Lab’de imitation/motion-tracking RL hedefi yap; "
    "sensör gürültüsü, gecikme, kütle, sürtünme ve motor parametrelerini domain randomization ile çeşitlendir. "
    "LLM, gerçek zamanlı eklem denetleyicisi yerine veri etiketleme ve görev düzeyi planlamada kullanılmalıdır.",
)
add_bullet(
    doc,
    "3 — Doğrulama ve sim2real: ",
    "Isaac Lab eğitiminden sonra politikayı MuJoCo’da bağımsız sim2sim test et; RViz’de TF, eklem, "
    "LiDAR/depth ve kalite topic’lerini izle; ancak ölçülebilir denge/limit testleri geçince askıda fiziksel G1’e aktar.",
)

add_heading(doc, "Sonuç", 1)
add_body(
    doc,
    "Proje bugün itibarıyla, resmi Stereolabs ve Unitree veri/model yapılarına dayanan, tekrar üretilebilir "
    "kayıtları kalite ölçümleriyle ayıran ve insan iskeletini fizik kurallarını koruyarak G1 23-DOF "
    "simülasyonlarına taşıyan işlevsel bir araştırma altyapısıdır. Sonraki kritik adım daha fazla rastgele "
    "kayıt toplamak değil; tek/çift kamera kalibrasyonu, etiketli hareket klipleri, güvenilir ayak teması ve "
    "denge politikası ile nicel doğrulamayı tamamlamaktır.",
    after=3,
)

refs = doc.add_paragraph()
set_paragraph_spacing(refs, before=3, after=0, line=1.0)
add_run(
    refs,
    "Resmî teknik dayanaklar: Stereolabs Body Tracking ve Fusion belgeleri; "
    "Unitree unitree_mujoco, unitree_rl_lab ve unitree_sim_isaaclab depoları; YanjieZe/GMR.",
    bold=True,
    color=DARK_BLUE,
    size=8.2,
)

core = doc.core_properties
core.title = "ZED 2i ile G1 EDU İnsan Hareketi Taklit Sistemi"
core.subject = "İki sayfalık teknik proje özeti"
core.author = "ZED G1 Projesi"
core.keywords = "ZED 2i, Unitree G1 EDU, BODY_38, GMR, MuJoCo, Isaac Lab, retargeting"

doc.save(OUT)
print(OUT)
